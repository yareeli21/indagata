# survey_intelligence/engine/readers/xlsx_reader.py
"""
Reader de XLSX (usa openpyxl).

Maneja los casos reales observados:
  - Un .xlsx es un ZIP+XML; openpyxl lo descomprime y parsea.
  - Valores datetime nativos (Google Forms 'Marca temporal') -> ISO string.
  - Selección de hoja (por nombre) o primera hoja por defecto.
  - Saltos de línea embebidos en encabezados (se normalizan).
"""
from __future__ import annotations

import datetime as _dt
import io

from survey_intelligence.engine.readers.header_utils import (
    normalize_cell,
    normalize_header,
)
from survey_intelligence.engine.readers.raw_table import RawTable


class XlsxReadError(Exception):
    """El contenido no pudo interpretarse como XLSX."""


def _cell_to_str(value: object) -> str:
    """Serializa una celda XLSX a str, normalizando fechas a ISO."""
    if value is None:
        return ""
    if isinstance(value, (_dt.datetime, _dt.date)):
        return value.isoformat()
    return normalize_cell(value)


def read_xlsx(data: bytes, sheet: str | None = None) -> RawTable:
    """
    Lee bytes XLSX y devuelve un RawTable.

    Args:
        data:  contenido del archivo .xlsx.
        sheet: nombre de hoja; None = primera hoja.

    Raises:
        XlsxReadError: si no puede abrirse o no tiene datos.
    """
    if not data:
        raise XlsxReadError("El archivo XLSX está vacío.")

    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise XlsxReadError("openpyxl no está disponible.") from exc

    try:
        wb = openpyxl.load_workbook(
            io.BytesIO(data), read_only=True, data_only=True
        )
    except Exception as exc:
        raise XlsxReadError(f"No se pudo abrir el XLSX: {exc}") from exc

    try:
        if sheet is not None:
            if sheet not in wb.sheetnames:
                raise XlsxReadError(
                    f"La hoja '{sheet}' no existe. Hojas: {wb.sheetnames}"
                )
            ws = wb[sheet]
        else:
            ws = wb[wb.sheetnames[0]]

        raw_rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()

    if not raw_rows:
        raise XlsxReadError("La hoja no contiene filas.")

    headers = [normalize_header(_cell_to_str(h)) for h in raw_rows[0]]
    n_cols = len(headers)

    data_rows: list[list[str]] = []
    for row in raw_rows[1:]:
        cells = [_cell_to_str(c) for c in row]
        if not any(c.strip() for c in cells):
            continue  # fila totalmente vacía
        if len(cells) < n_cols:
            cells += [""] * (n_cols - len(cells))
        elif len(cells) > n_cols:
            cells = cells[:n_cols]
        data_rows.append(cells)

    return RawTable(
        headers=headers,
        rows=data_rows,
        encoding=None,
        sheet=ws.title,
        notes=[],
    )
