"""Pruebas de archivos: integridad del crudo, versionado, lectura, borrado y permisos."""
from __future__ import annotations

import hashlib
import io
import json


def _subir(cliente, cabeceras, contenido: bytes, nombre: str, tipo: str = "raw", instrumento: str | None = None):
    datos = {"tipo": tipo}
    if instrumento:
        datos["instrumentoId"] = instrumento
    return cliente.post(
        "/almacenamiento/archivo",
        headers=cabeceras,
        files={"archivo": (nombre, contenido, "application/octet-stream")},
        data=datos,
    )


def test_el_crudo_se_guarda_sin_cambios_y_se_descarga_identico(cliente, servicio):
    contenido = "pregunta;respuesta\r\n¿Te gusta la escuela?;Sí\r\n".encode("cp1252")
    r = _subir(cliente, servicio, contenido, "Encuesta Ñandú.csv")
    assert r.status_code == 200, r.text
    ref = r.json()
    assert ref["sha256"] == hashlib.sha256(contenido).hexdigest()
    assert ref["key"].startswith("raw/_/1__")
    assert ref["nombre_original"] == "Encuesta Ñandú.csv"

    d = cliente.get("/almacenamiento/descarga", params={"ruta": ref["key"]}, headers=servicio)
    assert d.status_code == 200
    assert d.content == contenido
    assert d.headers["x-checksum-sha256"] == ref["sha256"]
    assert "filename*=UTF-8''Encuesta%20%C3%91and%C3%BA.csv" in d.headers["content-disposition"]

    # La ruta relativa (lo que se persiste en raw_data) también resuelve.
    d2 = cliente.get("/almacenamiento/descarga", params={"ruta": ref["ruta_relativa"]}, headers=servicio)
    assert d2.content == contenido


def test_versionado_por_instrumento(cliente, servicio):
    a = _subir(cliente, servicio, b"a,b\n1,2\n", "datos.csv", tipo="clean", instrumento="7").json()
    b = _subir(cliente, servicio, b"a,b\n3,4\n", "datos.csv", tipo="clean", instrumento="7").json()
    assert (a["version"], b["version"]) == (1, 2)
    lista = cliente.get("/almacenamiento/artefactos", params={"instrumentoId": "7"}, headers=servicio).json()
    assert [x["version"] for x in lista] == [1, 2]


def test_permisos(cliente, investigador, admin):
    assert _subir(cliente, {}, b"x", "a.csv").status_code == 401
    assert _subir(cliente, investigador, b"x", "a.csv").status_code == 403
    assert _subir(cliente, admin, b"x", "a.csv").status_code == 200


def test_json_ultima_version(cliente, servicio):
    for n in (1, 2):
        r = cliente.post(
            "/almacenamiento/json",
            headers=servicio,
            json={"instrumentoId": "12", "tipo": "metadata", "contenido": {"v": n}, "nombre": "Mi instrumento"},
        )
        assert r.status_code == 200, r.text
    detalle = cliente.get("/almacenamiento/artefacto/metadata/12", headers=servicio).json()
    assert detalle["version"] == 2
    assert detalle["contenido"] == {"v": 2}
    v1 = cliente.get("/almacenamiento/artefacto/metadata/12", params={"version": 1}, headers=servicio).json()
    assert v1["contenido"] == {"v": 1}
    assert cliente.get("/almacenamiento/artefacto/analysis/12", headers=servicio).status_code == 404


def test_documento_final_es_un_json_versionado_aparte(cliente, servicio):
    r = cliente.post("/almacenamiento/json", headers=servicio,
                     json={"instrumentoId": "12", "tipo": "documento", "contenido": {"esquema_version": "1.0"}, "nombre": "v1_documento_final"})
    assert r.status_code == 200, r.text
    ref = r.json()
    assert ref["key"].startswith("documento/12/1__") and ref["sha256"]
    detalle = cliente.get("/almacenamiento/artefacto/documento/12", headers=servicio).json()
    assert detalle["contenido"] == {"esquema_version": "1.0"}
    assert cliente.get("/almacenamiento/artefacto/metadata/12", headers=servicio).status_code == 404  # no se mezcla


def test_tabla_csv_punto_y_coma_latin1(cliente, servicio):
    contenido = "Carrera;Edad\nPedagogía;21\n;\nDerecho;22\n".encode("latin-1")
    ref = _subir(cliente, servicio, contenido, "r.csv").json()
    t = cliente.get("/almacenamiento/tabla", params={"ruta": ref["key"]}, headers=servicio).json()
    assert t["headers"] == ["Carrera", "Edad"]
    assert t["rows"] == [["Pedagogía", "21"], ["Derecho", "22"]]  # la fila vacía se omite
    assert t["n_filas"] == 2


def test_tabla_xlsx(cliente, servicio):
    import openpyxl

    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.append(["¿Qué tan satisfecho estás?", "Edad"])
    hoja.append([5, 20])
    hoja.append([4, 21.0])
    buffer = io.BytesIO()
    libro.save(buffer)
    ref = _subir(cliente, servicio, buffer.getvalue(), "respuestas.xlsx").json()
    t = cliente.get("/almacenamiento/tabla", params={"ruta": ref["key"]}, headers=servicio).json()
    assert t["rows"] == [["5", "20"], ["4", "21"]]


def test_tabla_desde_json_v1(cliente, servicio):
    data = {
        "instrument": {"sections": [{"questions": [{"question_id": "q1", "text": "¿Te gusta?"}]}]},
        "respondents": [
            {"respondent_id": "r1", "demographics": {"carrera": "Derecho"}, "answers": {"q1": 5}},
            {"respondent_id": "r2", "demographics": {"carrera": "Medicina"}, "answers": {"q1": None}},
        ],
    }
    ref = _subir(cliente, servicio, json.dumps(data).encode(), "inst.json").json()
    t = cliente.get("/almacenamiento/tabla", params={"ruta": ref["key"]}, headers=servicio).json()
    assert t["headers"] == ["respondent_id", "carrera", "¿Te gusta?"]
    assert t["rows"][1] == ["r2", "Medicina", ""]


def test_texto_docx_y_txt(cliente, servicio):
    import docx

    documento = docx.Document()
    documento.add_paragraph("Guía de entrevista")
    documento.add_paragraph("1. ¿Cómo describes tu experiencia?")
    buffer = io.BytesIO()
    documento.save(buffer)
    ref = _subir(cliente, servicio, buffer.getvalue(), "guia.docx").json()
    t = cliente.get("/almacenamiento/texto", params={"ruta": ref["key"]}, headers=servicio).json()
    assert "¿Cómo describes tu experiencia?" in t["texto"]

    ref_txt = _subir(cliente, servicio, "Hola\nmundo".encode("utf-8"), "notas.txt").json()
    t2 = cliente.get("/almacenamiento/texto", params={"ruta": ref_txt["key"]}, headers=servicio).json()
    assert t2["texto"] == "Hola\nmundo"


def test_ruta_heredada_del_seed_y_path_traversal(cliente, servicio):
    from services import file_manager

    destino = file_manager.raiz() / "raw" / "inst_01.v1.json"
    destino.write_text('{"instrument": {"title": "Demo"}}', encoding="utf-8")
    r = cliente.get("/almacenamiento/descarga", params={"ruta": "storage/raw/inst_01.v1.json"}, headers=servicio)
    assert r.status_code == 200
    assert r.json()["instrument"]["title"] == "Demo"

    malo = cliente.get("/almacenamiento/descarga", params={"ruta": "../../../etc/passwd"}, headers=servicio)
    assert malo.status_code in (400, 404)
    assert cliente.get("/almacenamiento/descarga", params={"ruta": "raw/no-existe.csv"}, headers=servicio).status_code == 404


def test_eliminar_y_purgar(cliente, servicio):
    ref = _subir(cliente, servicio, b"a\n1\n", "a.csv").json()
    r = cliente.delete("/almacenamiento/archivo", params={"ruta": ref["ruta_relativa"]}, headers=servicio)
    assert r.json() == {"ruta": ref["ruta_relativa"], "eliminado": True}
    assert cliente.get("/almacenamiento/descarga", params={"ruta": ref["key"]}, headers=servicio).status_code == 404
    # Idempotente.
    assert cliente.delete("/almacenamiento/archivo", params={"ruta": ref["key"]}, headers=servicio).json()["eliminado"] is False

    _subir(cliente, servicio, b"x", "x.csv", tipo="clean", instrumento="9")
    cliente.post("/almacenamiento/json", headers=servicio, json={"instrumentoId": "9", "tipo": "metadata", "contenido": {"a": 1}})
    p = cliente.delete("/almacenamiento/instrumentos/9", headers=servicio).json()
    assert p["artefactos_eliminados"] == 2
    assert cliente.get("/almacenamiento/artefactos", params={"instrumentoId": "9"}, headers=servicio).json() == []


def test_integridad_detecta_cambios(cliente, servicio):
    from services import file_manager

    ref = _subir(cliente, servicio, b"original", "a.txt").json()
    (file_manager.raiz() / ref["ruta_relativa"]).write_bytes(b"alterado")
    r = cliente.get("/almacenamiento/descarga", params={"ruta": ref["key"]}, headers=servicio)
    assert r.status_code == 500
    assert "huella" in r.json()["detail"]


def test_limpiar_temporales_y_estadisticas(cliente, servicio):
    from services import file_manager

    ref = _subir(cliente, servicio, b"tmp", "t.txt", tipo="temp").json()
    entradas = file_manager._leer_manifiesto()
    for e in entradas:
        if e["key"] == ref["key"]:
            e["creado"] = "2000-01-01T00:00:00+00:00"
    file_manager._escribir_manifiesto(entradas)
    r = cliente.post("/almacenamiento/mantenimiento/limpiar-temporales", headers=servicio)
    assert r.json() == {"eliminados": 1}
    est = cliente.get("/almacenamiento/estadisticas", headers=servicio).json()
    assert est["por_tipo"]["temp"]["n_artefactos"] == 0
