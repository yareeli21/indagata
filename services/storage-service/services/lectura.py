"""Lectura de archivos almacenados: tablas y texto plano.

Adaptado de services/instrument-service/app/parsing (los servicios no comparten
código de aplicación, solo `shared/`). El storage-service es quien conoce los
formatos en disco, así que expone estas lecturas a los demás servicios:

  - `leer_tabla`:  CSV / XLSX → encabezados + filas de texto. También convierte el
                   JSON de un instrumento (formato v1 de la demo) en tabla de
                   respuestas, para poder exportarlo a .sav.
  - `extraer_texto`: PDF / DOCX / TXT / MD / CSV / JSON → texto plano.

Las celdas siempre son `str` ("" si están vacías). Nunca se modifica el archivo.
"""
from __future__ import annotations

import csv
import datetime as _dt
import io
import json
import re
from dataclasses import dataclass, field
from typing import Any

_ENCODINGS = ("utf-8-sig", "utf-8", "cp1252", "latin-1")
_WS_RUN = re.compile(r"\s+")


class LecturaError(Exception):
    """El contenido no se pudo interpretar con el formato esperado."""


@dataclass
class Tabla:
    headers: list[str]
    rows: list[list[str]]
    encoding: str | None = None
    hoja: str | None = None
    notas: list[str] = field(default_factory=list)


@dataclass
class Documento:
    texto: str
    n_paginas: int | None = None
    notas: list[str] = field(default_factory=list)


# ── Normalización ──────────────────────────────────────────────────────────────


def _normalizar_encabezado(valor: object) -> str:
    if valor is None:
        return ""
    texto = str(valor).replace("\r", " ").replace("\n", " ").replace("\t", " ")
    return _WS_RUN.sub(" ", texto).strip()


def _normalizar_celda(valor: object) -> str:
    if valor is None:
        return ""
    if isinstance(valor, (_dt.datetime, _dt.date)):
        return valor.isoformat()
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return valor.strip() if isinstance(valor, str) else str(valor)


def _decodificar(data: bytes) -> tuple[str, str]:
    ultimo: Exception | None = None
    for enc in _ENCODINGS:
        try:
            return data.decode(enc), enc
        except UnicodeDecodeError as exc:
            ultimo = exc
    raise LecturaError(f"No se pudo decodificar el texto: {ultimo}")


def _ajustar(filas: list[list[str]], n_cols: int) -> list[list[str]]:
    """Rellena/recorta cada fila al nº de columnas y omite las totalmente vacías."""
    salida: list[list[str]] = []
    for fila in filas:
        if not any(c.strip() for c in fila):
            continue
        if len(fila) < n_cols:
            fila = fila + [""] * (n_cols - len(fila))
        salida.append(fila[:n_cols])
    return salida


# ── Tablas ─────────────────────────────────────────────────────────────────────


def _leer_csv(data: bytes) -> Tabla:
    texto, encoding = _decodificar(data)
    notas = [] if encoding in ("utf-8-sig", "utf-8") else [f"Codificación de respaldo: {encoding}"]
    try:
        delimitador = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t").delimiter
    except csv.Error:
        delimitador = ","
    filas = list(csv.reader(io.StringIO(texto), delimiter=delimitador))
    if not filas:
        raise LecturaError("El CSV no contiene filas.")
    headers = [_normalizar_encabezado(h) for h in filas[0]]
    cuerpo = [[_normalizar_celda(c) for c in fila] for fila in filas[1:]]
    return Tabla(headers=headers, rows=_ajustar(cuerpo, len(headers)), encoding=encoding, notas=notas)


def _leer_xlsx(data: bytes) -> Tabla:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise LecturaError("openpyxl no está disponible.") from exc
    try:
        libro = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise LecturaError(f"No se pudo abrir la hoja de cálculo: {exc}") from exc
    try:
        hoja = libro[libro.sheetnames[0]]
        filas = [list(r) for r in hoja.iter_rows(values_only=True)]
        titulo = hoja.title
    finally:
        libro.close()
    if not filas:
        raise LecturaError("La hoja no contiene filas.")
    headers = [_normalizar_encabezado(h) for h in filas[0]]
    cuerpo = [[_normalizar_celda(c) for c in fila] for fila in filas[1:]]
    return Tabla(headers=headers, rows=_ajustar(cuerpo, len(headers)), hoja=titulo)


def tabla_desde_json_instrumento(data: dict[str, Any]) -> Tabla:
    """Tabla de respuestas desde el JSON v1 (`respondents[].answers`).

    Columnas: datos demográficos + una por pregunta (encabezado = texto de la
    pregunta cuando está disponible en `instrument.sections[].questions[]`).
    """
    respondentes = [r for r in data.get("respondents") or [] if isinstance(r, dict)]
    if not respondentes:
        raise LecturaError("El JSON no contiene respuestas tabulares.")

    textos: dict[str, str] = {}
    for seccion in (data.get("instrument") or {}).get("sections") or []:
        for pregunta in (seccion or {}).get("questions") or []:
            if isinstance(pregunta, dict) and pregunta.get("question_id"):
                textos[str(pregunta["question_id"])] = _normalizar_encabezado(
                    pregunta.get("text") or pregunta["question_id"]
                )

    claves_demo: list[str] = []
    claves_resp: list[str] = []
    for r in respondentes:
        for clave in (r.get("demographics") or {}):
            if clave not in claves_demo:
                claves_demo.append(clave)
        for clave in (r.get("answers") or {}):
            if clave not in claves_resp:
                claves_resp.append(clave)

    headers = ["respondent_id", *claves_demo, *[textos.get(c, c) for c in claves_resp]]
    filas: list[list[str]] = []
    for r in respondentes:
        demo = r.get("demographics") or {}
        resp = r.get("answers") or {}
        filas.append(
            [_normalizar_celda(r.get("respondent_id"))]
            + [_normalizar_celda(demo.get(c)) for c in claves_demo]
            + [_normalizar_celda(resp.get(c)) for c in claves_resp]
        )
    return Tabla(headers=headers, rows=filas, notas=["Tabla generada desde el JSON del instrumento."])


def leer_tabla(data: bytes, extension: str) -> Tabla:
    """Lee un archivo tabular según su extensión."""
    if not data:
        raise LecturaError("El archivo está vacío.")
    ext = extension.lower()
    if ext == ".csv":
        return _leer_csv(data)
    if ext in (".xlsx", ".xls"):
        return _leer_xlsx(data)
    if ext == ".json":
        try:
            contenido = json.loads(_decodificar(data)[0])
        except json.JSONDecodeError as exc:
            raise LecturaError(f"JSON inválido: {exc}") from exc
        if not isinstance(contenido, dict):
            raise LecturaError("El JSON raíz debe ser un objeto.")
        return tabla_desde_json_instrumento(contenido)
    raise LecturaError(f"El formato {extension or '(sin extensión)'} no es tabular.")


# ── Texto ──────────────────────────────────────────────────────────────────────


def _texto_pdf(data: bytes) -> Documento:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise LecturaError("pypdf no está disponible.") from exc
    try:
        lector = PdfReader(io.BytesIO(data))
        if lector.is_encrypted:
            lector.decrypt("")
    except Exception as exc:
        raise LecturaError(f"No se pudo abrir el PDF: {exc}") from exc
    partes: list[str] = []
    for pagina in lector.pages:
        try:
            partes.append(pagina.extract_text() or "")
        except Exception:  # noqa: BLE001 - una página ilegible no aborta la lectura
            partes.append("")
    texto = "\n".join(partes).strip()
    notas = [] if texto else ["No se extrajo texto (posible PDF escaneado sin OCR)."]
    return Documento(texto=texto, n_paginas=len(lector.pages), notas=notas)


def _texto_docx(data: bytes) -> Documento:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise LecturaError("python-docx no está disponible.") from exc
    try:
        documento = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise LecturaError(f"No se pudo abrir el DOCX: {exc}") from exc
    lineas = [p.text for p in documento.paragraphs if p.text and p.text.strip()]
    for tabla in documento.tables:
        for fila in tabla.rows:
            celdas = [c.text.strip() for c in fila.cells if c.text and c.text.strip()]
            if celdas:
                lineas.append(" | ".join(celdas))
    return Documento(texto="\n".join(lineas).strip())


def _texto_json(data: bytes) -> Documento:
    """Texto legible del JSON de un instrumento (título, objetivo y preguntas)."""
    try:
        contenido = json.loads(_decodificar(data)[0])
    except json.JSONDecodeError as exc:
        raise LecturaError(f"JSON inválido: {exc}") from exc
    instrumento = contenido.get("instrument") if isinstance(contenido, dict) else None
    if not isinstance(instrumento, dict):
        return Documento(texto=json.dumps(contenido, ensure_ascii=False, indent=2))
    lineas: list[str] = []
    if instrumento.get("title"):
        lineas.append(f"# {instrumento['title']}")
    if instrumento.get("objective"):
        lineas.append(str(instrumento["objective"]))
    for seccion in instrumento.get("sections") or []:
        if not isinstance(seccion, dict):
            continue
        if seccion.get("name"):
            lineas.append(f"\n## {seccion['name']}")
        for n, pregunta in enumerate(seccion.get("questions") or [], start=1):
            if not isinstance(pregunta, dict):
                continue
            lineas.append(f"{n}. {pregunta.get('text', '')}")
            for opcion in pregunta.get("options") or []:
                lineas.append(f"   - {opcion}")
    if instrumento.get("contenido"):
        lineas.append(str(instrumento["contenido"]))
    return Documento(texto="\n".join(lineas).strip())


def extraer_texto(data: bytes, extension: str) -> Documento:
    """Extrae texto plano según la extensión."""
    if not data:
        raise LecturaError("El archivo está vacío.")
    ext = extension.lower()
    if ext == ".pdf":
        return _texto_pdf(data)
    if ext == ".docx":
        return _texto_docx(data)
    if ext in (".txt", ".md", ".csv"):
        texto, encoding = _decodificar(data)
        notas = [] if encoding in ("utf-8-sig", "utf-8") else [f"Codificación de respaldo: {encoding}"]
        return Documento(texto=texto.strip(), notas=notas)
    if ext == ".json":
        return _texto_json(data)
    if ext in (".xlsx", ".xls"):
        tabla = _leer_xlsx(data)
        return Documento(texto="\n".join(tabla.headers))
    raise LecturaError(f"No se puede extraer texto de archivos {extension or '(sin extensión)'}.")
